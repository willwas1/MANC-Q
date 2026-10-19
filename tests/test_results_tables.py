"""Checks of the results tables on the synthetic demo spectrum only (MANC-Q 1.8.2): upper bounds are never
stored as measured values, Not detected is 0 with the disclaimer, means use measured values and zeros only, and
regrade reproduces the tables. Fits EXPNO 10 of the demo once (about a minute or two).
Run with:  python -m pytest tests   (or:  python tests/test_results_tables.py)"""
import os
import sys
import tempfile
import numpy as np
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from mancq import engine as q
from mancq import demo as dm
from mancq.regrade import regrade

UB, ND = "Upper bound only", "Not detected"


def check_tables(res):
    G = pd.read_csv(os.path.join(res, "concentrations_long.csv"), dtype={"sample": str})
    ub = G[G.tier == UB]
    assert len(ub), "the demo should have some upper bounds"
    assert ub.conc_mM.isna().all(), "an upper bound was stored as a measured value"
    assert np.isfinite(ub.upper_bound_mM).all()
    assert (G[G.tier == ND].conc_mM == 0).all()
    xl = os.path.join(res, "metabolite_concentrations.xlsx")
    wb = openpyxl.load_workbook(xl, read_only=True)
    ws = wb["Concentration_mM_all"]
    assert ws["A1"].value == q.NOT_DETECTED_NOTE
    readme = " ".join(str(c.value) for row in wb["README"].iter_rows() for c in row if c.value)
    assert q.NOT_DETECTED_NOTE in readme
    wb.close()
    conc = pd.read_excel(xl, "Concentration_mM_all", header=3, index_col=0)
    report = pd.read_excel(xl, "Report_mM", index_col=0)
    summ = pd.read_excel(xl, "Summary", index_col=0)
    for r in ub.itertuples():
        col = int(r.sample) if int(r.sample) in conc.columns else r.sample
        assert np.isnan(conc.loc[r.metabolite, col])
        assert report.loc[r.metabolite, col] == f"<={r.upper_bound_mM:.2g}"
    assert {"n_measured", "n_not_detected"} <= set(summ.columns)
    for m in summ.index:
        v = G[(G.metabolite == m)].conc_mM
        if v.notna().any():
            assert abs(summ.loc[m, "mean_conc_mM"] - v.mean()) < 1e-9, m
    return G


def test_upper_bounds_not_values():
    with tempfile.TemporaryDirectory() as tmp:
        spec, res = os.path.join(tmp, "spectra"), os.path.join(tmp, "results")
        dm.write_demo(spec, cache_dir=os.path.join(tmp, "cache"))
        q.run(spec, res, dm.REF_MM, SAMPLES=["10"], OVERLAYS=False, N_WORKERS=1)
        a = check_tables(res)
        regrade(res)
        b = check_tables(res)
        assert (a.tier.values == b.tier.values).all()
        assert np.allclose(a.conc_mM.values, b.conc_mM.values, equal_nan=True, rtol=1e-5)


if __name__ == "__main__":
    test_upper_bounds_not_values()
    print("ok")
