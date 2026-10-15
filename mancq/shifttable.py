"""mancq.shifttable - the chemical shift assignment table for publication (supplementary table style).

One row per metabolite with a value in the chosen tiers (default Quantified, Overlapped or Manually adjusted) in
at least a set share of the samples (default 25 %, settings SHIFT_TABLE_TIERS, SHIFT_TABLE_MIN_FRACTION): its 1H
chemical shifts (fitted centre of each multiplet,
median over the samples where the metabolite was reported), with * on the peaks used to quantify it and ** on
peaks used but shared with another metabolite, how many samples it was reported in, and the reference spectrum
it was fitted with. Written as chemical_shift_assignments.xlsx (formatted, ready for a supplement) and as the
Shift_assignments sheet of the main workbook.
"""
import os
import re
import numpy as np
import pandas as pd

REPORTED = ("Quantified", "Overlapped (semi-quantitative)", "Manually adjusted")
MIN_FRACTION = 0.25
MIN_PEAK_SNR = 3.0          # a multiplet is listed if it was fitted at least this tall (x noise) in some sample
FILE = "chemical_shift_assignments.xlsx"


def display_name(name):
    """'L-Lactate' -> 'Lactate', 'Glutathione_reduced' -> 'Glutathione reduced'; 'myo-Inositol' stays."""
    n = re.sub(r"^(L|D)-", "", str(name)).replace("_", " ")
    return n[:1].upper() + n[1:] if not re.match(r"^(myo|cis|trans|beta|alpha|scyllo|N|O)-", n) else n


def reference_of(c):
    """Short reference-spectrum citation for a library compound."""
    prov = str(c.get("provenance", ""))
    if c.get("source") == "GISSMO":
        m = re.search(r"bmse\d+", prov)
        return f"GISSMO {m.group(0)}" if m else "GISSMO"
    if c.get("source") == "CASMDB":
        m = re.search(r"origin ([^)]*)\)", prov)
        return "CASMDB peak list" + (f" ({m.group(1)})" if m else "")
    if c.get("source") == "SPIN":
        return "literature spin system"
    return str(c.get("source", ""))


def build(G, P, lib, tiers=REPORTED, min_fraction=MIN_FRACTION):
    """DataFrame with columns Metabolite, 1H chemical shift / ppm, Reported in (samples), Reference spectrum,
    plus hidden helper columns (metabolite, n_marked). A metabolite is listed if it has a value in one of `tiers`
    in at least `min_fraction` of the samples (and in at least one)."""
    refs = {c["name"]: reference_of(c) for c in lib}
    n_samples = G["sample"].nunique()
    G = G[~G.metabolite.astype(str).str.startswith("Pluronic")]
    P = P.copy()
    P["sample"] = P["sample"].astype(str)
    rows = []
    for m, g in G.groupby("metabolite"):
        rep = g[g.tier.isin(tiers)]
        if not len(rep) or rep["sample"].nunique() < max(1, min_fraction * n_samples - 1e-9):
            continue
        samples = set(rep["sample"].astype(str))
        quant = set(rep[rep.tier == "Quantified"]["sample"].astype(str))
        over = set(rep[rep.tier == "Overlapped (semi-quantitative)"]["sample"].astype(str))
        pm = P[(P.metabolite == m) & P["sample"].isin(samples)]
        shifts = []
        for j, pj in pm.groupby("multiplet"):
            if not (pj.height_snr.astype(float) >= MIN_PEAK_SNR).any():
                continue                      # never fitted visibly (e.g. inside an ignored region)
            used = pj[pj.used_as_reporter.astype(bool)]["sample"]
            q, o = int(used.isin(quant).sum()), int(used.isin(over).sum())
            mark = ""
            if q + o >= max(1, 0.25 * len(samples)):
                mark = "*" if q >= o else "**"
            shifts.append((float(pj.centre_ppm.median()), mark))
        shifts.sort()
        text = ", ".join(f"{p:.3f}{k}" for p, k in shifts)
        rows.append({"Metabolite": display_name(m), "1H chemical shift / ppm": text,
                     "Reported in (samples)": f"{len(samples)} of {n_samples}",
                     "Reference spectrum": refs.get(m, ""), "metabolite": m,
                     "n_marked": sum(1 for _, k in shifts if k)})
    df = pd.DataFrame(rows, columns=["Metabolite", "1H chemical shift / ppm", "Reported in (samples)",
                                     "Reference spectrum", "metabolite", "n_marked"])
    return df.sort_values("Metabolite", key=lambda s: s.str.lower()).reset_index(drop=True)


def footnote(n_spectra, sf=None, version="", reference="TSP", tiers=REPORTED, min_fraction=MIN_FRACTION):
    field = f" at {sf:.0f} MHz" if sf else ""
    names = [t.split(" (")[0].lower() for t in tiers]
    which = ", ".join(names[:-1]) + (" or " if len(names) > 1 else "") + names[-1]
    share = f"in at least {min_fraction:.0%} of the spectra" if min_fraction > 0 else "in at least one spectrum"
    return (f"Listed: metabolites {which} {share}. 1H chemical shifts (ppm) are referenced to {reference} at 0.000 ppm. Each is the centre of a fitted "
            f"multiplet, the median over the samples in which the metabolite was reported (of {n_spectra} spectra "
            f"recorded{field}, fitted with MANC-Q {version}). Shifts marked * are the peaks used to quantify the "
            f"metabolite; shifts marked ** were used but are shared with another metabolite, so those values are "
            f"semi-quantitative. Metabolites with no marked shift were estimated from the fit of their whole "
            f"spectrum. pH-sensitive metabolites (for example histidine and citrate) move between samples; the "
            f"position in each sample is in peak_list_by_metabolite.csv. Reference spectra: GISSMO (BMRB) spin "
            f"systems simulated at the spectrometer field, or CASMDB peak lists.")


def write(path, df, note, title="Chemical shift assignments for metabolites"):
    """The table as a formatted Excel file: title, ruled header, rows, footnote (journal supplementary style)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Border, Side, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Chemical shifts"
    cols = ["Metabolite", "1H chemical shift / ppm", "Reported in (samples)", "Reference spectrum"]
    rule = Side(style="thin", color="000000")
    f = Font(name="Calibri", size=10)
    ws.cell(row=1, column=1, value=title).font = Font(name="Calibri", size=11, bold=True)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=3, column=j, value=c)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.border = Border(top=rule, bottom=rule)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=3, column=2).value = "\u00b9H chemical shift / ppm"
    for i, r in enumerate(df[cols].itertuples(index=False), 4):
        for j, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = f
            cell.alignment = Alignment(vertical="top", wrap_text=(j == 2))
    last = 3 + len(df)
    for j in range(1, len(cols) + 1):
        ws.cell(row=last, column=j).border = Border(bottom=rule, top=rule if last == 3 else None)
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=len(cols))
    c = ws.cell(row=last + 2, column=1, value=note)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[last + 2].height = 15 * (len(note) // 110 + 2)
    for col, w in zip("ABCD", (24, 70, 20, 30)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    wb.save(path)


def make(out, G, P, Q, lib, version="", reference="TSP", tiers=REPORTED, min_fraction=MIN_FRACTION):
    """Build the table and write chemical_shift_assignments.xlsx; returns (DataFrame, footnote)."""
    df = build(G, P, lib, tiers, min_fraction)
    sf = float(np.nanmedian(Q.SF_MHz)) if "SF_MHz" in Q and len(Q) else None
    note = footnote(G["sample"].nunique(), sf, version, reference, tiers, min_fraction)
    write(os.path.join(out, FILE), df, note)
    return df, note
